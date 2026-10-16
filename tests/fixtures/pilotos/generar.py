# -*- coding: utf-8 -*-
"""
LOS DOS PILOTOS FICTICIOS (instrucción de Dirección, 13-10-2026, sección 19).

Genera los archivos sintéticos Word, PDF y Excel de «Modas Lucero» (tienda de ropa) y «Estudio Vargas &
Asociados» (jurídico). Todo es inventado: empresas, personas, cifras. Los archivos llevan FALTANTES y
CONTRADICCIONES deliberadas, listadas en declaraciones.json, para probar que el sistema las ve y las dice,
no para probar efectividad comercial.

Determinista: misma semilla → mismos archivos. Reproducir: python tests/fixtures/pilotos/generar.py
Requiere python-docx, openpyxl, reportlab (instalados en la máquina de desarrollo; no son dependencias del repo).
"""
import json, os, random
from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

AQUI = os.path.dirname(os.path.abspath(__file__))
random.seed(20261013)


def pdf(path, titulo, lineas):
    c = canvas.Canvas(path, pagesize=A4)
    ancho, alto = A4
    y = alto - 60
    c.setFont("Helvetica-Bold", 14); c.drawString(50, y, titulo); y -= 30
    c.setFont("Helvetica", 10)
    for ln in lineas:
        if y < 60:
            c.showPage(); c.setFont("Helvetica", 10); y = alto - 60
        c.drawString(50, y, ln); y -= 15
    c.save()


# ───────────────────────────── PILOTO A · MODAS LUCERO ─────────────────────────────
A = os.path.join(AQUI, "modas-lucero"); os.makedirs(A, exist_ok=True)

# Word · Funciones del personal (INCOMPLETO: Jorge sin funciones; Carmen sin horario; dice que Rosa fija ofertas)
d = Document()
d.add_heading("MODAS LUCERO S.A.C. — Funciones del personal", 0)
d.add_paragraph("Versión de marzo de 2025. Elaborado por la gerencia.")
for puesto, persona, funciones, horario in [
    ("Gerente general", "Lucero Quintana", ["Compra de mercadería en Gamarra (semanal)", "Fija precios de toda la tienda", "Cierre de caja y depósito", "Pago a proveedores", "Publicaciones en Instagram y TikTok", "Atención en mostrador"], "L–S 10:00–20:00"),
    ("Vendedora principal / caja", "Rosa Mamani", ["Atención y cobro en caja", "Reposición de percheros", "Fija precios de las ofertas de fin de temporada", "Cuadre de caja al cierre junto con la gerente"], "L–S 10:00–19:00"),
    ("Vendedora", "Carmen Huamán", ["Atención al cliente", "Probadores y doblado", "Apoyo en inventario cuando se pide"], ""),
    ("Almacén y reparto (medio tiempo)", "Jorge Ccahuana", [], "L–V 14:00–18:00"),
]:
    d.add_heading(f"{puesto} — {persona}", level=2)
    if funciones:
        for f in funciones: d.add_paragraph(f, style="List Bullet")
    else:
        d.add_paragraph("Funciones: POR DEFINIR.")
    d.add_paragraph(f"Horario: {horario}" if horario else "Horario: (no indicado)")
d.add_paragraph("Nota: no hay manual de caja ni de inventario. Las devoluciones las autoriza la gerente por WhatsApp.")
d.save(os.path.join(A, "funciones-del-personal.docx"))

# PDF · Contrato de alquiler
pdf(os.path.join(A, "contrato-alquiler.pdf"), "CONTRATO DE ARRENDAMIENTO DE LOCAL COMERCIAL", [
    "Arrendador: Inmobiliaria San Borja E.I.R.L. (ficticia)   Arrendataria: Modas Lucero S.A.C. (ficticia)",
    "Inmueble: tienda comercial de 48 m2, Av. Los Pinos 234, Santiago de Surco, Lima.",
    "Plazo: 24 meses, del 01 de marzo de 2025 al 28 de febrero de 2027.",
    "Renta mensual: S/ 6,500.00 (seis mil quinientos y 00/100 soles), pagadera hasta el dia 5 de cada mes.",
    "Garantia: dos (2) meses de renta, S/ 13,000.00, entregados a la firma.",
    "Servicios: luz, agua e internet a cargo de la arrendataria. Mantenimiento de galeria: S/ 180.00 mensuales.",
    "Penalidad por mora: 0.5% diario sobre la renta impaga.",
    "Lima, 20 de febrero de 2025.",
])

# Excel · Ventas, compras, gastos (INCOMPLETO) e inventario (40 de 120 sin costo)
wb = Workbook(); ws = wb.active; ws.title = "Ventas"
ws.append(["mes", "ventas_contado", "ventas_tarjeta", "total_ventas", "compras_mercaderia", "devoluciones"])
ventas = [("2026-04", 74200, 0.42, 48900, 1850), ("2026-05", 81500, 0.45, 55200, 2100), ("2026-06", 69800, 0.40, 44100, 1600),
          ("2026-07", 88300, 0.47, 58600, 2400), ("2026-08", 79100, 0.44, 51300, 1900), ("2026-09", 86900, 0.46, 54700, 2300)]
for mes, total, pct_tarj, compras, dev in ventas:
    tarj = round(total * pct_tarj); ws.append([mes, total - tarj, tarj, total, compras, dev])
ws2 = wb.create_sheet("Gastos")
ws2.append(["mes", "alquiler", "planilla", "luz_agua_internet", "mantenimiento_galeria", "transporte", "comision_tarjeta", "otros"])
# FALTANTE deliberado: julio y agosto sin planilla ni luz; nadie anotó transporte en 3 meses
ws2.append(["2026-04", 6500, 9000, 480, 180, 350, 1040, 220])
ws2.append(["2026-05", 6500, 9000, 510, 180, None, 1150, 300])
ws2.append(["2026-06", 6500, 9000, 470, 180, 380, 920, None])
ws2.append(["2026-07", 6500, None, None, 180, None, 1300, 250])
ws2.append(["2026-08", 6500, None, None, 180, 410, 1100, 180])
ws2.append(["2026-09", 6500, 9000, 495, 180, None, 1250, 270])
ws3 = wb.create_sheet("Inventario")
ws3.append(["codigo", "descripcion", "categoria", "talla", "stock", "costo_unitario", "precio_venta"])
cats = [("Blusa", 38, 89), ("Pantalón", 52, 119), ("Vestido", 61, 149), ("Polo", 22, 49), ("Casaca", 95, 219), ("Falda", 34, 79), ("Short", 27, 59), ("Chompa", 58, 139)]
total_costo = 0
for i in range(1, 121):
    nombre, costo, precio = cats[i % len(cats)]
    stock = random.randint(8, 30)
    sin_costo = (i % 3 == 0)  # 40 de 120 sin costo: documento incompleto
    c = None if sin_costo else costo + random.randint(-4, 6)
    if c: total_costo += c * stock
    ws3.append([f"ML-{i:03d}", f"{nombre} modelo {i}", nombre, random.choice(["S", "M", "L", "XL"]), stock, c, precio + random.choice([0, 0, 10, -10])])
ws3.append([]); ws3.append(["TOTAL A COSTO (solo los que tienen costo)", None, None, None, None, total_costo, None])
wb.save(os.path.join(A, "ventas-e-inventario.xlsx"))

# ───────────────────────────── PILOTO B · ESTUDIO VARGAS ─────────────────────────────
B = os.path.join(AQUI, "estudio-vargas"); os.makedirs(B, exist_ok=True)

pdf(os.path.join(B, "tarifario-2025.pdf"), "ESTUDIO VARGAS & ASOCIADOS — TARIFARIO DE HONORARIOS 2025", [
    "Vigente desde enero de 2025 (ficticio).",
    "Laboral — demanda de beneficios sociales o despido: S/ 1,800.00 por caso, hasta sentencia de primera instancia.",
    "Civil — obligacion de dar suma de dinero / desalojo: S/ 2,500.00 por caso.",
    "Familia — alimentos, tenencia, regimen de visitas: S/ 1,200.00 por caso.",
    "Consulta presencial o virtual: S/ 150.00 (se descuenta si el cliente contrata).",
    "Los honorarios no incluyen tasas judiciales ni gastos de notificacion.",
    "Forma de pago: 50% a la firma, 50% al admitirse la demanda.",
])

d = Document()
d.add_heading("Procedimiento de seguimiento de casos (borrador)", 0)
for p in [
    "Todo expediente nuevo lo abre el Dr. Vargas y define la estrategia.",
    "La Dra. Paredes prepara los escritos; el Dr. Vargas los revisa y firma todos.",
    "El Sr. Quispe lleva los cargos al juzgado y avisa por el grupo de WhatsApp del estudio.",
    "Los vencimientos se anotan en la agenda personal del Dr. Vargas. No hay tablero compartido.",
    "Al cliente se le informa cuando pregunta. No hay informe periódico.",
    "Cobranza: el Dr. Vargas recuerda al cliente por llamada cuando se acuerda.",
]:
    d.add_paragraph(p, style="List Bullet")
d.add_paragraph("Última revisión: no consta.")
d.save(os.path.join(B, "procedimiento-seguimiento.docx"))

wb = Workbook(); ws = wb.active; ws.title = "Expedientes 2026-I"
ws.append(["codigo", "cliente", "materia", "abogado_responsable", "fecha_inicio", "estado", "honorario_pactado", "cobrado", "horas_estimadas", "ultima_actuacion"])
nombres = ["Flores", "Torres", "Ramos", "Castro", "Vega", "Rojas", "Medina", "Chávez", "Salazar", "Paredes", "Ñahui", "Quispe", "Huanca", "Condori", "Mamani"]
materias = [("Laboral", 62, 1800, (1200, 2200)), ("Civil", 30, 2500, (2000, 2800)), ("Familia", 18, 1200, (900, 1400))]
n = 0
for materia, cuantos, tarifa, rango in materias:
    for _ in range(cuantos):
        n += 1
        pactado = random.choice([tarifa, tarifa, random.randint(*rango)])  # CONTRADICCIÓN: descuentos no escritos
        cobrado = random.choice([pactado, pactado // 2, None, pactado])  # FALTANTE: cobros sin registrar
        ws.append([f"EXP-{n:03d}", f"{random.choice(nombres)} {random.choice(nombres)}", materia,
                   random.choice(["Dr. Vargas", "Dr. Vargas", "Dra. Paredes"]) if materia != "Civil" else "Dr. Vargas",
                   f"2026-0{random.randint(1, 6)}-{random.randint(1, 28):02d}",
                   random.choice(["en trámite", "en trámite", "con sentencia", "archivado"]), pactado, cobrado,
                   None,  # DESCONOCIDO deliberado: nadie registra horas
                   f"2026-{random.randint(7, 9):02d}-{random.randint(1, 28):02d}"])
wb.save(os.path.join(B, "expedientes-2026-I.xlsx"))

# ───────────────────────────── DECLARACIONES Y CONTRADICCIONES ─────────────────────────────
declaraciones = {
    "modas_lucero": {
        "empresa": "Modas Lucero S.A.C. (ficticia)", "sector": "Comercio minorista de ropa", "ciudad": "Lima (Surco)", "antiguedad_anios": 5, "personas": 4,
        "entrada_breve": {
            "negocio_hoy": "Tienda de ropa de mujer en Surco, 5 años, somos 4: yo, Rosa, Carmen y Jorge medio tiempo.",
            "que_vende_estrella": "Blusas y vestidos; lo que más deja son los vestidos de fiesta.",
            "ventas_mes": 80000, "ganancia_mes": 8000, "meta_venta_mes": 160000,
            "deuda_clientes": "Casi nada, vendemos al contado y con tarjeta.",
            "vision_3_anios": "Dos tiendas, una más en Miraflores, y yo ya no en el mostrador.",
            "rol_futuro": "Comprar, decidir colecciones y ver los números. Ya no cobrar ni doblar ropa.",
            "dinero_tranquila": 15000, "no_sacrifica": "No despido a nadie; Rosa lleva cinco años conmigo. Y los domingos no trabajo.",
            "que_se_detiene_sin_mi": "Las compras y los precios. Rosa puede cobrar y cerrar caja; Carmen no.",
            "decision_postergada": "Contratar a alguien que me reemplace en tienda. Me cuesta cada mes seguir en el mostrador 10 horas.",
            "principal_problema": "No sé cuánto gano de verdad; el inventario está desordenado.",
            "que_funciona": "La clientela vuelve; Instagram trae gente nueva.",
            "recursos_desaprovechados": "Tengo stock parado de la temporada pasada y un almacén que no se usa bien.",
            "mejor_resultado_anterior": "Diciembre 2025: vendimos 140,000 en el mes con la campaña de fiestas.",
            "stock_declarado": "Tengo como 40,000 en mercadería.",
        },
        "archivos": ["funciones-del-personal.docx", "contrato-alquiler.pdf", "ventas-e-inventario.xlsx"],
        "contradicciones_deliberadas": [
            "Stock declarado «como 40,000» vs Excel Inventario: total a costo de los que tienen costo supera 70,000 (y 40 artículos sin costo).",
            "Declara que fija TODOS los precios vs Word: Rosa fija los precios de las ofertas de fin de temporada.",
            "Utilidad declarada 8,000 vs Excel: ventas ~80,000 − compras ~52,000 − alquiler 6,500 − planilla 9,000 − otros conocidos ≈ 10,000–12,000 ANTES de los gastos que faltan (transporte, luz de dos meses, mermas, impuestos): la cifra no se puede confirmar ni negar; hay que preguntar.",
        ],
        "faltantes_deliberados": [
            "Inventario: 40 de 120 artículos sin costo unitario → margen por producto no calculable para ellos.",
            "Gastos: planilla y luz sin registrar en julio y agosto; transporte sin registrar en tres meses.",
            "Word: Jorge sin funciones; Carmen sin horario. Sin manual de caja ni de inventario.",
            "Ningún registro de clientes → recompra «de cada 10 vuelven 4» no es medible hoy.",
        ],
    },
    "estudio_vargas": {
        "empresa": "Estudio Vargas & Asociados (ficticio)", "sector": "Servicios jurídicos", "ciudad": "Lima (Jesús María)", "antiguedad_anios": 9, "personas": 3,
        "entrada_breve": {
            "negocio_hoy": "Estudio jurídico, 9 años, tres personas: yo (socio), la Dra. Paredes y el Sr. Quispe, asistente.",
            "que_vende_estrella": "Laboral, civil y familia. En número, laboral; en honorarios, civil. No sé cuál deja más.",
            "ventas_mes": "Entre 18,000 y 25,000 según el mes; no lo tengo exacto.", "ganancia_mes": "No lo sé.", "meta_venta_mes": None,
            "deuda_clientes": "Varios clientes deben la segunda mitad; no tengo la suma.",
            "vision_3_anios": "Que el estudio funcione si yo me enfermo un mes. Que nos busquen por laboral, no por baratos.",
            "rol_futuro": "Estrategia de los casos grandes y relación con clientes. Que la Dra. Paredes lleve laboral de principio a fin.",
            "dinero_tranquila": 12000, "no_sacrifica": "No bajo honorarios para competir. No tomo casos penales.",
            "que_se_detiene_sin_mi": "Todo: firmo todos los escritos, llevo la agenda de vencimientos en mi cuaderno y cobro yo.",
            "decision_postergada": "Contratar otro abogado o darle laboral completo a la Dra. Paredes. Me cuesta que ella no pueda responder al cliente sin preguntarme.",
            "principal_problema": "No sé qué servicio me deja plata; el seguimiento es informal y se nos han pasado dos plazos este año.",
            "que_funciona": "Ganamos la mayoría de los laborales; los clientes nos recomiendan.",
            "recursos_desaprovechados": "La Dra. Paredes podría llevar más; el asistente tiene tiempo libre por las tardes.",
            "mejor_resultado_anterior": "2024: un caso civil grande que dejó 18,000 en honorarios.",
            "horas_por_caso_laboral": "Unas 25 horas por caso, calculo.",
        },
        "archivos": ["tarifario-2025.pdf", "procedimiento-seguimiento.docx", "expedientes-2026-I.xlsx"],
        "contradicciones_deliberadas": [
            "Tarifario laboral 1,800 vs Excel: honorarios pactados entre 1,200 y 2,200 → descuentos no escritos.",
            "«Unas 25 horas por caso» × 62 laborales en un semestre ≈ 1,550 horas para dos abogados → ~30 h/semana solo laboral, más civil y familia: la estimación o la carga no cuadran; se pregunta, no se concluye.",
            "Procedimiento dice que el Dr. Vargas lleva vencimientos en su agenda vs declaración «se nos han pasado dos plazos».",
        ],
        "faltantes_deliberados": [
            "Excel: columna horas_estimadas vacía en todos los expedientes → rentabilidad por servicio DESCONOCIDA (resultado válido, se dice).",
            "Excel: cobrado vacío en parte de los expedientes → cobranza pendiente no sumable.",
            "Tarifario 2025: vigencia no confirmada en 2026.",
            "Sin registro de tiempo de respuesta al cliente → estándar «48 horas» no medible hoy.",
        ],
    },
    "casos_dificiles": {
        "informacion_desconocida": "Vargas: ganancia_mes «No lo sé», horas por expediente vacías → el sistema debe guardar «desconocido», no 0.",
        "documento_incompleto": "Lucero: Inventario con 40 artículos sin costo; Gastos con celdas vacías → se dice qué falta, nunca un promedio inventado.",
        "contradiccion": "Lucero: stock 40,000 declarado vs >70,000 a costo en Excel → ambas cifras quedan declarado/extraído pendientes.",
        "falla_proveedor": "Proveedor falso (src/lib/ai/simulador.ts) lanzando AIProviderPausedError (402) en medio de la extracción.",
        "reintento_misma_operacion": "El mismo POST de subida/extracción dos veces → un solo trabajo (clave de idempotencia determinista).",
        "usuario_sin_acceso": "El dueño de Vargas pide /api/findings?company_id=<Lucero> → 403 y rastro en auditoría.",
        "cambio_version_linea_base": "Lucero: línea base v1 con utilidad 8,000 declarada; corrección a v2 tras costos completos; los cortes anteriores siguen refiriendo v1.",
    },
}
with open(os.path.join(AQUI, "declaraciones.json"), "w", encoding="utf-8") as f:
    json.dump(declaraciones, f, ensure_ascii=False, indent=1)
print("generados:", sorted(os.listdir(A)), sorted(os.listdir(B)), "declaraciones.json")
